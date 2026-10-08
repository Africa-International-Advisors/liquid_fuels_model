"""The extended demand baseline workbook: sheets present, inputs carried faithfully, Nigel's cells untouched."""
import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from lfm.scripts import build_demand_baseline_workbook as build

ROOT = Path(__file__).resolve().parents[1]
VINTAGE = ROOT / "assumptions" / "2026"
pytestmark = pytest.mark.skipif(not (ROOT / build.OUT).exists(), reason="workbook not built")


def _rows(ws):
    return {row[0].value: row for row in ws.iter_rows(min_row=5) if isinstance(row[0].value, str)}


def test_history_carries_the_registered_inputs_and_uses_formulas_for_the_balance():
    d = build.load(VINTAGE / "timeseries", VINTAGE / "reference")
    history = _rows(load_workbook(ROOT / build.OUT)["History"])
    at = {year: build.FIRST - 1 + build.YEARS.index(year) for year in build.YEARS}

    assert history["Imports, diesel"][at[2024]].value == pytest.approx(d["sars"][("import", "diesel", 2024)])
    assert history["Imports, diesel"][at[2013]].value is None          # customs is not in litres before 2014
    assert history["Department national file, petrol"][at[2023]].value == pytest.approx(d["dept"][("petrol", 2023)])
    assert history["Department national file, petrol"][at[2024]].value is None
    assert history["Gauteng diesel"][at[2022]].value == pytest.approx(d["province"][("diesel", "GP", 2022)])
    assert history["Production reported, diesel"][at[2022]].value is None    # no energy balance after 2021
    for label in ("Sum of nine provinces, diesel", "Net imports, diesel", "Sales used, diesel",
                  "Sales less net imports, diesel", "Litres levied less recorded sales"):
        assert str(history[label][at[2024]].value).startswith("="), label


def test_only_the_source_selection_changes_in_nigels_sheets():
    before, after = load_workbook(ROOT / build.SOURCE), load_workbook(ROOT / build.OUT)
    assert [ws.title for ws in after.worksheets] == [ws.title for ws in before.worksheets] + [
        "History", "DR01 balance", "Sector history", "Diesel by use", "Demand by use", "DR04 routes", "DR04 entry points", "DR04 transport cost", "DR07 evidence", "DR07 power diesel",
        "DR07 fleet by province", "DR07 efficiency and rail", "Power fleet", "Vehicle history", "HML response", "Gap status",
        "Checks",
        "History sources"]
    changed = []
    for ws in before.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                new = after[ws.title][cell.coordinate].value
                # Saving rewrites floats to fifteen significant digits; that is not a change.
                same_number = (isinstance(new, (int, float)) and isinstance(cell.value, (int, float))
                               and new == pytest.approx(cell.value, rel=1e-12))
                if new != cell.value and not same_number:
                    changed.append((ws.title, cell.coordinate, cell.value, new))
    assert changed and all(sheet == "Source selection" and (old, new) == ("FIASA", "Department")
                           for sheet, _, old, new in changed)


def test_every_checked_evidence_value_matches_the_inputs():
    results = [row[5].value for row in load_workbook(ROOT / build.OUT)["Checks"].iter_rows(min_row=5)]
    assert results.count("matches") > 50 and "differs" not in results


def test_provincial_estimates_are_marked_and_kept_apart_from_observations():
    from lfm.scripts import backtest_provincial_shares as shares

    history = _rows(load_workbook(ROOT / build.OUT)["History"])
    at = {year: build.FIRST - 1 + build.YEARS.index(year) for year in build.YEARS}
    brought_in = history["Gauteng petrol"][at[2023]]                         # estimate, linked from section 6
    assert str(brought_in.value).startswith("=IF(ISNUMBER(") and brought_in.font.italic
    assert brought_in.fill.fgColor.rgb == build.FILL["estimate"]
    assert history["Gauteng petrol"][at[2022]].fill.fgColor.rgb == build.FILL["observation"]
    assert "ESTIMATES" in history["Gauteng petrol"][4].value
    assert history["Gauteng petrol"][at[2025]].value.startswith("=")          # blank until a national figure exists
    for product in ("petrol", "diesel"):
        assert "FIASA is not used" in history[f"Why 2024 and 2025 are blank, {product}"][4].value
        assert "FIASA" not in str(history[f"Sales used, {product}"][at[2024]].value)      # department only
    estimated = [history[f"{name} share of petrol"][at[2023]].value for name in build.PROVINCE_NAMES.values()]
    assert sum(estimated) == pytest.approx(100)
    quarter1 = shares.share(shares.load(VINTAGE / "timeseries")["quarter1"][("petrol", 2023)])
    assert history["Gauteng share of petrol"][at[2023]].value == pytest.approx(quarter1["GP"] * 100)
    assert history["Gauteng share of petrol"][2].value == "Proposed estimate"
    assert str(history["Gauteng petrol, estimated"][at[2024]].value).startswith("=")
    assert history["Gauteng petrol, estimated"][at[2022]].value is None


def test_quarter_one_shares_beat_holding_last_year_in_the_back_test():
    from lfm.scripts import backtest_provincial_shares as shares

    scores = {(r["product"], r["method"], r["years_ahead"]): r["mean_share_points_misallocated"]
              for r in shares.backtest(shares.load(VINTAGE / "timeseries"))}
    for product in ("petrol", "diesel"):
        assert scores[(product, "quarter1", 0)] < scores[(product, "hold", 1)] < scores[(product, "trend3", 3)]


def test_vehicle_sheet_carries_observed_stock_and_the_models_settings():
    d = build.load(VINTAGE / "timeseries", VINTAGE / "reference")
    vehicles = _rows(load_workbook(ROOT / build.OUT)["Vehicle history"])
    at = {year: build.FIRST - 1 + build.YEARS.index(year) for year in build.YEARS}
    assert vehicles["Cars"][at[2025]].value == d["stock"][("cars", 2025)]
    assert vehicles["Cars, new sales"][at[2024]].value == d["new_sales"][("cars", 2024)]
    assert vehicles["Diesel vehicles registered"][at[2023]].value == d["by_fuel_2023"]["diesel"]
    assert vehicles["Cars retirement rate"][3].value == pytest.approx(d["model_vehicles"]["scrappage_rate"]["passenger"] * 100)
    assert vehicles["Cars share of new sales"][3].value == pytest.approx(78)
    for label in ("Cars retired", "Cars retirement rate", "Implied less registered diesel vehicles",
                  "Distance implied by petrol sales"):
        assert str(vehicles[label][at[2023]].value).startswith("="), label


def test_diesel_by_use_splits_road_diesel_with_the_studys_shares():
    d = build.load(VINTAGE / "timeseries", VINTAGE / "reference")
    sheet = _rows(load_workbook(ROOT / build.OUT)["Diesel by use"])
    at = {year: build.FIRST - 1 + build.YEARS.index(year) for year in build.YEARS}
    shares = [sheet[group][3].value for group in build.VEHICLE_GROUPS]
    assert sum(shares) == pytest.approx(100)
    assert sheet["Heavy vehicles"][3].value == pytest.approx(60.5, abs=0.1)
    assert set(d["study_classes"]["Heavy vehicles"]) == {f"HCV{n}Diesel" for n in range(1, 10)}
    assert d["study_classes"]["Light vehicles"] == ["LCVDiesel"]
    assert "MBTDiesel" in d["study_classes"]["Passenger vehicles"] and "BusDiesel" in d["study_classes"]["Passenger vehicles"]
    assert "HCV1Diesel" in sheet["Heavy vehicles: classes"][4].value
    assert sheet["Fuel bought by road freight businesses"][at[2023]].value == pytest.approx(d["freight_floor"][2023][0])
    for label in ("Road vehicles and uses not listed above", "Heavy vehicles", "Sum of the six branches"):
        assert str(sheet[label][at[2023]].value).startswith("="), label


def test_lever_response_gap_status_jet_and_sources_are_in_the_workbook():
    wb = load_workbook(ROOT / build.OUT)
    levers = [row for row in wb["HML response"].iter_rows(min_row=5, values_only=True) if row[0]]
    assert len(levers) == 29                                                 # 20 proposed levers and 9 added
    results = [row[9] for row in levers]
    assert results.count("added lever") == 9 and "replacement proposed" in results and "no change" in results
    gaps = [row[0] for row in wb["Gap status"].iter_rows(min_row=5, values_only=True) if row[0]]
    assert gaps == [f"G{n:02d}" for n in range(1, 14)]
    history = _rows(wb["History"])
    at = {year: build.FIRST - 1 + build.YEARS.index(year) for year in build.YEARS}
    assert history["Department national file, jet"][at[2023]].value == pytest.approx(1843.97, abs=0.01)
    assert str(history["Jet sold per aircraft movement"][at[2024]].value).startswith("=")
    used_for = " ".join(str(row[0]) for row in wb["History sources"].iter_rows(min_row=5, values_only=True))
    for sheet in ("History section 7", "Diesel by use", "Vehicle history", "HML response"):
        assert sheet in used_for, sheet


def test_power_fleet_lists_the_diesel_stations_and_builds_three_cases():
    rows = _rows(load_workbook(ROOT / build.OUT)["Power fleet"])
    first = build.FIRST - 1                                                  # 2022; +8 is 2030, +9 is 2031
    assert [rows[name][first].value for name in ("Ankerlig (Eskom)", "Gourikwa (Eskom)", "Avon (Independent producer)",
                                                 "Dedisa (Independent producer)")] == [1338, 746, 670, 335]
    assert rows["Acacia (Eskom)"][2].value == "Not counted"                  # kerosene, not diesel
    assert rows["Camden"][first + 7].value == 1561 and rows["Camden"][first + 8].value == 0
    low, medium = rows["Avon and Dedisa available, low"], rows["Avon and Dedisa available, medium"]
    assert [low[first + n].value for n in (7, 8, 9)] == [1005, 0, 0]         # both contracts end in 2030 (system operator)
    assert [medium[first + n].value for n in (7, 8, 9, 13)] == [1005] * 4    # medium: both continue
    assert [rows[f"Gas turbines at the five coal sites, {case}"][first + 8].value
            for case in ("low", "medium", "high")] == [0, 3000, 6000]
    assert rows["Share of Ankerlig and Gourikwa output on diesel, high"][first + 8].value == "=100"
    assert "Camden, Grootvlei, Hendrina, Arnot and Kriel" in rows["High case"][4].value
    for label in ("Diesel for power, low case", "Diesel for power, high case", "High case",
                  "Ceiling: four stations and 6 GW of gas turbines on diesel all year"):
        assert str(rows[label][first + 8].value).startswith("="), label


def test_power_case_totals_follow_the_three_stories():
    cases = build.power_case_totals(build.load(VINTAGE / "timeseries", VINTAGE / "reference"))
    assert cases["low"][2035] < cases["medium"][2035] < cases["high"][2035] < cases["ceiling"][2035]
    assert cases["medium"][2028] < cases["medium"][2027]                      # Ankerlig and Gourikwa switch to gas
    assert cases["low"][2031] < cases["low"][2029]                            # Avon and Dedisa gone
    assert cases["medium"][2030] > cases["medium"][2029]                      # gas turbines at coal sites, diesel backup
    assert cases["high"][2029] == pytest.approx(cases["high"][2025])          # no gas switch in the high case
    assert cases["high"][2030] == pytest.approx(4223.5, abs=1)                # value shown by the recalculated workbook
    assert cases["medium"][2030] == pytest.approx(637.0, abs=0.5)
    assert cases["low"][2035] == pytest.approx(20.3, abs=0.1)


def test_every_row_on_the_added_sheets_names_a_source():
    wb = load_workbook(ROOT / build.OUT)
    for name in ("History", "Sector history", "Diesel by use", "Power fleet", "Vehicle history"):
        ws = wb[name]
        head = next(r for r in ws.iter_rows(min_row=1, max_row=8) if "Source" in [c.value for c in r])
        at = [c.value for c in head].index("Source")
        missing = [r[0].value for r in ws.iter_rows(min_row=head[0].row + 1)
                   if r[0].value and r[2].value and r[2].value != "Note" and not r[at].value]
        assert missing == [], (name, missing[:5])
    history = _rows(wb["History"])
    assert "Commodity-Flow-and-Energy-Balance" in history["Production reported, diesel"][at].value
    assert history["Refinery output reported to JODI, diesel"][2].value == "Comparison only"
    assert "Production used, diesel" not in history


def test_dr01_sheet_is_formulas_on_history_in_billion_litres_with_a_source_on_every_line():
    ws = load_workbook(ROOT / build.OUT)["DR01 balance"]
    header = [c.value for c in ws[4]]
    assert header[0] == "Line" and header[1:13] == list(range(2014, 2026)) and header[13:] == ["Source", "Note"]
    lines = {}
    block = None
    for r in ws.iter_rows(min_row=5):
        if r[0].value in ("Diesel", "Petrol", "Diesel and petrol"):
            block = r[0].value
        elif r[0].value:
            lines[(block, r[0].value)] = r
    assert [k[1] for k in lines if k[0] == "Diesel"] == ["Sales", "Production", "Imports", "Exports", "Supply",
                                                         "Stock change", "Supply less sales"]
    assert all(r[13].value for r in lines.values())                                # a source on every line
    sales_2021 = lines[("Diesel", "Sales")][1 + build.DR01_YEARS.index(2021)].value
    assert sales_2021.startswith("=IF(ISNUMBER(History!") and sales_2021.endswith('/1000,"")')
    assert all(c.value is None for c in lines[("Petrol", "Stock change")][1:13])    # not published


def test_demand_by_use_sheet_lists_each_use_in_billion_litres_with_basis_and_source():
    ws = load_workbook(ROOT / build.OUT)["Demand by use"]
    header = [c.value for c in ws[4]]
    assert header[0] == "Use" and header[1:11] == list(range(2014, 2024)) and header[11:] == ["Basis", "Source", "Note"]
    labels = [r[0].value for r in ws.iter_rows(min_row=5) if r[0].value]
    assert labels[:9] == ["Diesel", "Power generation", "Mining", "Manufacturing and other industry", "Agriculture",
                          "Heavy vehicles", "Light vehicles", "Passenger vehicles", "Diesel sales"]
    assert "Petrol" in labels and labels[-1] == "Petrol sales"
    for r in ws.iter_rows(min_row=6):
        if r[0].value and r[0].value not in ("Diesel", "Petrol"):
            assert r[11].value and r[12].value, r[0].value                      # basis and source on every line
    mining = next(r for r in ws.iter_rows(min_row=5) if r[0].value == "Mining")
    assert mining[1].value.startswith("=IF(ISNUMBER('Diesel by use'!") and mining[1].value.endswith('/1000,"")')


def test_dr04_routes_sheet_groups_every_fact_with_status_source_and_gap():
    ws = load_workbook(ROOT / build.OUT)["DR04 routes"]
    assert [c.value for c in ws[4]] == ["Item", "Asset or route", "Value", "Unit", "Period", "Status", "Source", "Page", "Open gap"]
    facts = [r for r in ws.iter_rows(min_row=5) if r[5].value]
    with (ROOT / build.DR04_EVIDENCE).open(encoding="utf-8", newline="") as fh:
        evidence = list(csv.DictReader(fh))
    kept = [r for r in evidence if r["scope"] != "other products included"]
    assert len(facts) == len(kept) < len(evidence)                                     # mixed-product facts are left out
    shown = {r[0].value for r in facts}
    assert not shown & {"Liquid bulk landed (imports)", "Liquid bulk landed, by month", "Petroleum volumes transported"}
    assert "Pipeline tariff, Durban to Alrode" in shown and "Diesel imported by road from Mozambique" in shown
    assert all(r[2].value and r[6].value for r in facts)                               # a value (or "Not available") and a source
    sections = [r[0].value for r in ws.iter_rows(min_row=5) if r[0].value and not r[5].value]
    assert sections[:2] == ["Pipeline limit", "Pipeline cost"] and {"Access", "Competing routes"} <= set(sections)
    assert "Pipeline use" not in sections                       # those volumes include crude and jet, so the section is empty
    open_rows = [r for r in facts if r[5].value not in ("observed", "inferred")]
    assert open_rows and all(r[0].fill.fgColor.rgb == build.FILL["estimate"] for r in open_rows)


def test_dr04_entry_points_sheet_is_petrol_and_diesel_only_and_adds_to_national_imports():
    wb = load_workbook(ROOT / build.OUT)
    assert "DR04 ports" not in wb.sheetnames                                           # all-liquids tonnage is not shown
    ws = wb["DR04 entry points"]
    header = [c.value for c in ws[4]]
    assert header[:2] == ["Entry point", "Route"] and header[2:14] == list(range(2014, 2026)) and header[14:] == ["Source", "Note"]
    blocks, block = {}, None
    for r in ws.iter_rows(min_row=5):
        if r[0].value in ("Diesel", "Petrol", "Diesel and petrol"):
            block = r[0].value
        elif r[0].value:
            blocks[(block, r[0].value)] = r
    at = 2 + build.ENTRY_YEARS.index(2025)
    d = build.load(VINTAGE / "timeseries", VINTAGE / "reference")
    for product, name in (("diesel", "Diesel"), ("petrol", "Petrol")):
        offices = [k[1] for k in blocks if k[0] == name and k[1] not in ("All offices", "Durban share")]
        total = sum(blocks[(name, o)][at].value for o in offices)
        assert total == pytest.approx(d["sars"][("import", product, 2025)] / 1000, rel=1e-6)   # adds to national customs imports
        assert blocks[(name, "All offices")][at].value.startswith("=SUM(")
        assert all(blocks[(name, o)][14].value for o in offices)                                # a source on every row
    assert blocks[("Diesel", "Komatipoort")][1].value == "Road, from Mozambique"


def test_dr04_transport_cost_sheet_lists_every_zone_with_a_source_and_marks_unpublished_rates():
    ws = load_workbook(ROOT / build.OUT)["DR04 transport cost"]
    elements = {r[0].value: r for r in ws.iter_rows(min_row=6, max_row=11)}
    assert elements["Pipeline tariff, Durban to Alrode"][1].value == 67.99
    assert elements["Regulated transport differential, Gauteng (zone 9C)"][3].value == 91.1
    assert elements["Commercial road tanker rate, Durban to Gauteng"][1].value is None
    assert elements["Rail rate, Durban to Gauteng"][4].value == "Not published"
    assert elements["Rail rate, Durban to Gauteng"][0].fill.fgColor.rgb == build.FILL["estimate"]       # open lines are yellow
    assert all(r[7].value for r in elements.values())
    assert [c.value for c in ws[14]] == ["Zone", "Provinces", "Districts", "Examples", "2014", "2024", "Change", "Source"]
    zones = {r[0].value: r for r in ws.iter_rows(min_row=15) if r[0].value}
    assert len(zones) == 54
    assert (zones["9C (Gauteng)"][4].value, zones["9C (Gauteng)"][5].value) == (33.1, 82.8)
    assert zones["1A (coast)"][5].value == 3.8 and "Durban" in zones["1A (coast)"][3].value
    assert zones["9C (Gauteng)"][6].value.startswith("=IF(COUNT(") and all(r[7].value for r in zones.values())


def test_dr07_sheets_show_reported_eskom_litres_the_implied_rate_and_the_fleet_by_province():
    wb = load_workbook(ROOT / build.OUT)
    evidence = wb["DR07 evidence"]
    facts = [r for r in evidence.iter_rows(min_row=5) if r[5].value]
    assert len(facts) >= 25 and all(r[2].value and r[6].value for r in facts)
    missing = [r[0].value for r in facts if r[5].value == "not available"]
    assert "Diesel burned in private backup generators" in missing and "Vehicles by age" in missing
    assert "Fuel use of new light vehicles, by year" in {r[0].value for r in facts if r[5].value == "observed"}
    assert not any("FIASA" in str(r[6].value) or "JODI" in str(r[6].value) for r in facts)

    power = {r[0].value: r for r in wb["DR07 power diesel"].iter_rows(min_row=5) if r[0].value}
    at = 1 + build.POWER_YEARS.index(2024)
    assert power["Fuel burned, as reported"][at].value == pytest.approx(1.1295)                # 1,129.5 million litres, Eskom p.141
    assert power["Electricity generated by Eskom, GWh"][at].value == 3634
    assert power["Electricity generated by independent plants, GWh"][at].value == 1509
    assert power["Litres per kWh, implied"][at].value.startswith("=IF(COUNT(")
    assert 1129.5 / 3634 == pytest.approx(build.LITRES_PER_KWH, abs=0.002)                     # the rate used is the rate reported
    assert power["Fuel burned, as reported"][1 + build.POWER_YEARS.index(2026)].value is None  # not yet reported
    assert all(r[12].value for r in power.values() if r[1].value is not None or r[at].value is not None)

    fleet = {r[0].value: r for r in wb["DR07 fleet by province"].iter_rows(min_row=5) if r[0].value}
    assert list(fleet)[:3] == ["Gauteng", "KwaZulu-Natal", "Western Cape"] and "South Africa" in fleet
    d = build.load(VINTAGE / "timeseries", VINTAGE / "reference")
    assert fleet["Gauteng"][1].value == d["natis_by_province"][("GP", "cars")]
    national = sum(fleet[name][1].value for _, name in build.FLEET_PROVINCES)
    assert national == pytest.approx(d["natis_by_province"][("ZAF", "cars")], rel=1e-9)        # provinces add to the national register


def test_dr07_efficiency_and_rail_sheet_carries_both_histories_with_sources():
    ws = load_workbook(ROOT / build.OUT)["DR07 efficiency and rail"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=5) if r[0].value}
    years = [c.value for c in rows["Year"][1:10]]
    assert years == [2005, 2008, 2010, 2011, 2012, 2013, 2014, 2015, 2019]
    assert [c.value for c in rows["Litres of petrol equivalent per 100 km"][1:10]] == [8.8, 8.6, 8.6, 8.2, 7.9, 7.7, 7.8, 7.8, 7.4]
    assert (7.4 / 8.8) ** (1 / 14) - 1 == pytest.approx(-0.012, abs=0.001)                  # close to the IEA's stated 1.3% a year
    rail = [c.value for c in rows["Million tonnes"][1:10]]
    assert rail == [226.3, 215.1, 212.4, 183.3, 173.1, 149.5, 151.7, 160.1, 250.0]
    assert rows["Year to March"][9].value == "Target"
    for label in ("Litres of petrol equivalent per 100 km", "Million tonnes"):
        assert "IEA" in rows[label][-2].value or "Transnet" in rows[label][-2].value
