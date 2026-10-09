"""The data requests workbook: one sheet a request, yearly figures in year columns, and a sheet of what was not found."""
import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from lfm.scripts import build_data_requests_workbook as requests

ROOT = Path(__file__).resolve().parents[1]
VINTAGE = ROOT / "assumptions" / "2026"
pytestmark = pytest.mark.skipif(not (ROOT / requests.OUT).exists(), reason="workbook not built")


def test_workbook_has_one_main_sheet_a_request_with_its_detail_tables_and_a_contents_page():
    wb = load_workbook(ROOT / requests.OUT)
    assert wb.sheetnames == [
        "Contents", "DR01 national balance", "DR01 sales by province", "DR01 demand by sector", "DR04 routes and access", "DR04 entry points",
        "DR04 transport cost", "DR07 demand evidence", "DR07 fleet by province", "DR08 refinery supply", "Market sizing", "Not available"]
    listed = [r[0].value for r in wb["Contents"].iter_rows(min_row=5, max_row=4 + len(wb.sheetnames) - 1)]
    assert listed == wb.sheetnames[1:]                               # every sheet is described on the contents page


def test_series_are_split_into_period_columns_and_breakdowns_into_rows():
    row = {"value": "458.7; 580.4; 937.5", "period": "years to March 2021; 2022; 2023", "asset_or_route": "Eskom", "unit": "million litres"}
    assert requests.as_series(row) == ("years to March", ["2021", "2022", "2023"], [458.7, 580.4, 937.5])
    assert requests.as_series({**row, "value": "Shut down; planned; delayed"}) is None           # words are not a series
    stations = {"value": "1,338; 746", "period": "current", "asset_or_route": "Ankerlig; Gourikwa", "unit": "MW"}
    assert requests.as_breakdown(stations) == [("Ankerlig", 1338), ("Gourikwa", 746)]
    assert requests.as_breakdown({**stations, "value": "cars 8,340,592; trucks 409,049", "asset_or_route": "South Africa"}) == [
        ("cars", 8340592), ("trucks", 409049)]
    assert requests.as_breakdown({**stations, "value": "Two weeks for a vessel; two days for road loading"}) is None


def test_no_cell_on_a_request_sheet_packs_several_figures_together():
    wb = load_workbook(ROOT / requests.OUT)
    for title, *_ in requests.REQUESTS:
        packed = [c.coordinate for row in wb[title].iter_rows(min_row=4) for c in row[4:12]
                  if isinstance(c.value, str) and c.value.count(";") and all(requests._NUMBER.match(v.strip()) for v in c.value.split(";"))]
        assert packed == [], (title, packed[:5])


def test_request_sheets_carry_every_fact_of_the_tables_in_year_columns():
    wb = load_workbook(ROOT / requests.OUT)
    for title, request, path, parts, _, _ in requests.REQUESTS:
        rows = requests.evidence(ROOT / path, VINTAGE / "reference", request)
        items = {r[0].value for r in wb[title].iter_rows(min_row=4) if r[0].value}
        assert {r["item"] for r in rows} <= items, title
        grouped = [r["asset_or_route"] for r in rows if r["item"] in ("Coal station shutdown", "Kerosene stations: shutdown", "Independent diesel plants: contract end",
                                                                      "Output while not operating", "Terminal operator at Island View")
                   and (" and " in r["asset_or_route"] or "," in r["asset_or_route"]) and "about 22 GW" not in r["asset_or_route"]]
        assert grouped == [], (title, grouped)                        # one station, plant or operator to a row
        assert set(parts) & items                                     # the section headings
    ws = wb["DR07 demand evidence"]
    header = next(r for r in ws.iter_rows(min_row=4) if r[0].value == "Item")
    years = [c.value for c in header[4:14] if isinstance(c.value, int)]
    fuel = next(r for r in ws.iter_rows(min_row=4) if r[0].value and r[0].value.startswith("Fuel burned at Eskom"))
    assert fuel[4 + years.index(2024)].value == pytest.approx(1129.5) and fuel[3].value == "years to March"
    assert fuel[4 + years.index(2016)].value == pytest.approx(1247.8)  # all ten years Eskom reports, not the last five
    stations = [r[1].value for r in ws.iter_rows(min_row=4) if r[0].value == "Diesel stations and capacity"]
    assert stations == ["Ankerlig", "Gourikwa", "Avon", "Dedisa"]     # one row a station
    sources = [str(c.value) for title, *_ in requests.REQUESTS for r in wb[title].iter_rows(min_row=4) for c in r[-3:-2]]
    assert not any(s.startswith("FIASA") or "JODI" in s for s in sources)


def test_dr04_shows_the_gauteng_transport_cost_by_year_from_the_departments_tables():
    rows = requests.gauteng_transport_rows(VINTAGE / "reference")
    diesel = next(r for r in rows if "diesel" in r["item"])
    series = dict(zip(diesel["period"].replace("year end ", "").split("; "), diesel["value"].split("; ")))
    assert (series["2012"], series["2014"], series["2023"], series["2024"]) == ("26.8", "33.1", "75.7", "75.7")
    assert list(series) == [str(y) for y in range(2012, 2025)] and "Department" in diesel["source"]
    ws = load_workbook(ROOT / requests.OUT)["DR04 routes and access"]
    shown = next(r for r in ws.iter_rows(min_row=4) if r[0].value == diesel["item"])
    assert [c.value for c in shown[4:17]] == [float(v) for v in diesel["value"].split("; ")]


def test_dr01_balance_is_in_billion_litres_with_a_source_on_every_line_and_blanks_where_unpublished():
    ws = load_workbook(ROOT / requests.OUT)["DR01 national balance"]
    lines, block = {}, None
    for r in ws.iter_rows(min_row=4):
        if r[0].value in ("Petrol", "Diesel"):
            block = r[0].value
        elif r[0].value and r[0].value != "Line":
            lines[(block, r[0].value)] = r
    assert [k[1] for k in lines if k[0] == "Diesel"] == ["Sales", "Production", "Imports", "Exports", "Stock change", "Supply less sales"]
    at = {year: 1 + requests.BALANCE_YEARS.index(year) for year in requests.BALANCE_YEARS}
    with (ROOT / requests.BALANCE).open(encoding="utf-8", newline="") as fh:
        balance = {(r["product"], int(r["period"])): r for r in csv.DictReader(fh)}
    assert lines[("Diesel", "Sales")][at[2021]].value == pytest.approx(float(balance[("diesel", 2021)]["sales_used"]) / 1e9)
    assert lines[("Diesel", "Production")][at[2022]].value is None and lines[("Petrol", "Sales")][at[2024]].value is None
    assert str(lines[("Diesel", "Supply less sales")][at[2021]].value).startswith("=IF(COUNT(")
    assert all(r[13].value for r in lines.values())                    # a source on every line


def test_sales_by_province_adds_to_the_national_sales_line_and_marks_2023_as_an_estimate():
    wb = load_workbook(ROOT / requests.OUT)
    ws = wb["DR01 sales by province"]
    blocks, block = {}, None
    for r in ws.iter_rows(min_row=4):
        if r[0].value in ("Petrol", "Diesel"):
            block = r[0].value
        elif r[0].value and r[0].value != "Province" and block:
            blocks.setdefault(block, {})[r[0].value] = r
    diesel = blocks["Diesel"]
    assert list(diesel)[:9] == ["Gauteng", "KwaZulu-Natal", "Western Cape", "Eastern Cape", "Mpumalanga", "Free State", "North West", "Limpopo",
                                "Northern Cape"]
    at = {year: 1 + requests.PROVINCE_YEARS.index(year) for year in requests.PROVINCE_YEARS}
    provinces = list(diesel.values())[:9]
    with (ROOT / requests.BALANCE).open(encoding="utf-8", newline="") as fh:
        balance = {(r["product"], int(r["period"])): float(r["sales_used"]) / 1e9 for r in csv.DictReader(fh) if r["sales_used"]}
    for year in (2015, 2021, 2023):                                    # the provinces add to the sales line of the national balance
        assert sum(r[at[year]].value for r in provinces) == pytest.approx(balance[("diesel", year)], rel=1e-6)
    fills = {r[at[2023]].fill.fgColor.rgb for r in provinces}, {r[at[2022]].fill.fgColor.rgb for r in provinces}
    assert fills == ({requests.FILL["estimate"]}, {requests.FILL["observation"]})
    assert all(r[at[2024]].value is None and r[at[2025]].value is None for r in provinces)
    assert all(r[14].value for r in list(diesel.values())[:11])        # a source on every line: nine provinces, the total, the national file


def test_demand_by_sector_carries_the_energy_balance_figures_for_mining_industry_and_agriculture():
    ws = load_workbook(ROOT / requests.OUT)["DR01 demand by sector"]
    blocks, block = {}, None
    for r in ws.iter_rows(min_row=4):
        if r[0].value in ("Petrol", "Diesel"):
            block = r[0].value
        elif r[0].value and r[0].value != "Sector":
            blocks.setdefault(block, {})[r[0].value] = r
    diesel = blocks["Diesel"]
    assert {"Mining", "Manufacturing and other industry", "Agriculture", "Road transport", "All sectors"} <= set(diesel)
    at = {year: 1 + requests.SECTOR_YEARS.index(year) for year in requests.SECTOR_YEARS}
    with (VINTAGE / "timeseries" / "energy_balance_department.csv").open(encoding="utf-8", newline="") as fh:
        balance = {(r["product"], r["flow_key"], int(r["period"])): float(r["value"]) / 1e9 for r in csv.DictReader(fh)}
    assert diesel["Mining"][at[2021]].value == pytest.approx(balance[("diesel", "mining", 2021)])
    assert diesel["Agriculture"][at[2014]].value == pytest.approx(balance[("diesel", "agriculture", 2014)])
    assert all(r[9].value for r in diesel.values())                    # a source on every line


def test_dr08_says_which_plants_are_refining_and_holds_the_assumed_astron_split():
    from lfm.scripts import build_dr08_refinery_evidence as dr08

    rows = dr08.build(dr08.load(VINTAGE))
    refining = {r["asset_or_route"]: r["value"] for r in rows if r["item"] == "Refining in 2026"}
    assert refining == {"Secunda": "Yes", "Natref": "Yes", "Astron Energy": "Yes", "Sapref": "No", "Enref": "No", "PetroSA": "No"}
    split = {r["item"]: r for r in rows if r["asset_or_route"] == "Astron Energy" and "assumed" in r["item"]}
    assert {r["value"] for r in split.values()} == {"50"} and len(split) == 2
    assert all(r["status"] == "inferred" and "Assumption" in r["source"] for r in split.values())
    idle = {r["asset_or_route"] for r in rows if r["item"] == "Output while not operating"}
    assert idle == {"Sapref", "Enref", "PetroSA"}


def test_not_available_sheet_lists_every_open_item_with_its_reason():
    rows = requests.missing_rows(VINTAGE / "reference")
    ws = load_workbook(ROOT / requests.OUT)["Not available"]
    shown = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=5) if r[0].value]
    assert shown == rows and all(why for *_, why in rows)
    what = {(request, item) for request, item, *_ in rows}
    assert {("DR01", "Production after 2021"), ("DR01", "Sales for 2024 and 2025"), ("DR07", "Diesel burned in private backup generators"),
            ("DR04", "Commercial road tanker rate, Durban to Gauteng"), ("DR08", "Petrol and diesel output by plant")} <= what
    assert not any("Ceiling" in item for _, item, *_ in rows)
