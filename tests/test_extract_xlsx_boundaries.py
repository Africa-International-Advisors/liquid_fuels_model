"""The workbook extractor must stop at the end of the table it was sent to read."""

import csv

import openpyxl

from lfm.scripts import extract_xlsx_to_assumptions as extract


def test_leading_year_rows_stops_at_the_first_row_without_a_year() -> None:
    rows = [
        [2022, 0.9], [2023, 0.8],
        [None, None],
        ["Low production scenario", None],
        [2022, 0.5],          # a second table further down the same range
    ]
    assert extract.leading_year_rows(rows) == [[2022, 0.9], [2023, 0.8]]


def test_historical_demand_ignores_tables_below_the_national_one(tmp_path, monkeypatch) -> None:
    wb = openpyxl.Workbook()
    for name in ("Gasoline - DemandSupply", "Diesel - DemandSupply"):
        ws = wb.create_sheet(name)
        ws["C11"], ws["K11"] = 2019, 100.0
    jet = wb.create_sheet("Jet - DemandSupply")
    jet["E9"], jet["J9"] = 2019, 2449.0
    jet["E10"], jet["J10"] = "2020*", 1098.0     # annotated year still counts
    jet["E12"] = "Jet Fuel"                       # regional table starts here
    jet["E13"], jet["J13"] = 2019, 382.0
    jet["E14"], jet["J14"] = 2020, 122.0

    monkeypatch.setattr(extract, "OUT", tmp_path)
    monkeypatch.setattr(extract, "REPO", tmp_path)
    extract.extract_historical_demand(wb)

    with (tmp_path / "historical_demand.csv").open(encoding="utf-8") as fh:
        rows = [r for r in csv.DictReader(fh) if r["product"] == "jet_a1"]
    assert [(r["period"], float(r["value"])) for r in rows] == [("2019", 2449.0), ("2020", 1098.0)]
