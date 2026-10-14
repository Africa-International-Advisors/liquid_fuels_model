"""Slide 13 for 2025: market of interest (six provinces) and its import requirement (Henry Gilfillan, 13 October 2026).

2022 district sales (Department, latest full year) give each district's share of national sales by product; those shares are applied
to 2025 national demand by product (production + imports - exports, p4). That is the 2022 data grown at national growth, on p4's basis.
Usage: python s13_2025.py PACK16.pptx OUT.pptx REPO_ROOT      (the new slide is appended as the last slide of OUT)
"""
import collections, copy, csv, sys
from pathlib import Path
import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC, OUT, REPO = sys.argv[1], sys.argv[2], Path(sys.argv[3])
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
T, A = "assumptions/2026/timeseries/", "output/delivered/sa_analytics_2026_10_11_v02/"
PD = ("petrol", "diesel")
prs = Presentation(SRC)
assert len(prs.slides) == 16


def rows(rel):
    with open(REPO / rel, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


# ---- 2025 national demand and plant output: same calculation as slides 2, 4 and 10 (s_demand.py) ----
ops = rows("assumptions/2026/reference/refinery_output_operators.csv")
sas = {(r["plant"], r["period"]): float(r["value"]) for r in ops if r["operator"] == "Sasol"}
ast = {r["period"]: float(r["value"]) for r in ops if r["operator"] == "Glencore"}
plant = {r["Plant"]: r for r in rows(A + "plant_supply.csv")}
yv = lambda p, k: float(plant[p][k])
BBL, BTU = 158.987294928 / 1000, yv("Astron", "All products bn L") / ast["2025"]
allp = {"Secunda": sas["Secunda", "2025"] * BBL, "Natref": sas["Natref", "2025"] / 0.6364 * BBL, "Astron": ast["2025"] * BTU}
pp = {k: {"petrol": v * yv(k, "Petrol yield"), "diesel": v * yv(k, "Diesel yield")} for k, v in allp.items()}
tr = rows(T + "fuel_trade_sars.csv")


def trade(flow, p):
    rr = [r for r in tr if r["period"] == "2025" and r["flow"] == flow and r["product"] == p]
    assert rr and all(r["unit"] == "litres" and r["months_reported"] == "12" for r in rr), (flow, p)
    return sum(float(r["value"]) for r in rr) / 1e9


nat25 = {p: sum(pp[k][p] for k in pp) + trade("import", p) - trade("export", p) for p in PD}
assert abs(sum(nat25.values()) - 25.38) < 0.01

# ---- district sales, 2022: four quarterly sheets; a province row is followed by its districts ----
CODE = {"Eastern Cape": "EC", "Free State": "FS", "Gauteng": "GP", "KwaZulu Natal": "KZN", "Limpopo": "LP", "Mpumalanga": "MP", "North West": "NW",
        "Northern Cape": "NC", "Western Cape": "WC"}


def province_of(name):
    n = name.replace("-", " ").replace(" Province", "").replace(" ", "").strip().lower()
    for k, c in CODE.items():
        if n == k.replace(" ", "").lower():
            return c
    return None


wb = openpyxl.load_workbook(REPO / "data/raw/energy_dept/fsv_district/2022-Quarter4-Magisterial-Districts-data.xlsx", read_only=True, data_only=True)
dist = collections.defaultdict(lambda: [0.0, 0.0])
ptot = collections.defaultdict(lambda: [0.0, 0.0])
f = lambda v: float(v) if v not in (None, "") else 0.0
for sn in ("2022 Q1", "2022 Q2", "2022 Q3", "2022 Q4"):
    hdr, prov = None, None
    for row in wb[sn].iter_rows(values_only=True):
        if row[0] == "Row Labels":
            hdr = [str(c) if c else "" for c in row]
            ip, idd = hdr.index("Petrol"), hdr.index("Diesel")
            continue
        if hdr is None or row[0] is None:
            continue
        name = str(row[0]).strip()
        if name.lower().startswith("grand total"):
            continue
        p = province_of(name)
        if p:
            prov = p
            ptot[p][0] += f(row[ip]) / 1e9
            ptot[p][1] += f(row[idd]) / 1e9
        else:
            assert prov, name
            dist[prov, name][0] += f(row[ip]) / 1e9
            dist[prov, name][1] += f(row[idd]) / 1e9
assert set(ptot) == set(CODE.values()), sorted(ptot)
for c in CODE.values():
    assert abs(sum(v[0] + v[1] for (q, _), v in dist.items() if q == c) - sum(ptot[c])) < 1e-6, c
n22 = [sum(v[0] for v in ptot.values()), sum(v[1] for v in ptot.values())]          # national 2022 sales by product
g = [nat25["petrol"] / n22[0], nat25["diesel"] / n22[1]]                           # 2022 sales -> 2025 demand, by product
print("national 2022 sales", [round(x, 2) for x in n22], "| 2025 demand", {k: round(v, 2) for k, v in nat25.items()}, "| factors", [round(x, 3) for x in g])
esc = lambda v: [v[0] * g[0], v[1] * g[1]]
reg = {c: esc(v) for c, v in ptot.items()}
d25 = {k: esc(v) for k, v in dist.items()}
MOI = [("GP", "Gauteng"), ("KZN", "KwaZulu-Natal"), ("MP", "Mpumalanga"), ("FS", "Free State"), ("NW", "North West"), ("LP", "Limpopo")]
codes = [c for c, _ in MOI]
nat = sum(nat25.values())
mp, md = sum(reg[c][0] for c in codes), sum(reg[c][1] for c in codes)
mt = mp + md
mt22 = sum(sum(ptot[c]) for c in codes)
md_list = sorted(((n, q, v) for (q, n), v in d25.items() if q in codes), key=lambda x: -(x[2][0] + x[2][1]))
TOP = md_list[:8]
n_d = len(md_list)
ins_p = pp["Secunda"]["petrol"] + pp["Natref"]["petrol"]
ins_d = pp["Secunda"]["diesel"] + pp["Natref"]["diesel"]
inside = ins_p + ins_d
need_p, need_d = mp - ins_p, md - ins_d
need = need_p + need_d
print(f"MOI 2025 {mt:.2f} ({mt / nat:.0%}) petrol {mp:.2f} diesel {md:.2f} | 2022 {mt22:.2f} | inside {inside:.2f} (P {ins_p:.2f} D {ins_d:.2f}) | "
      f"requirement {need:.2f} (P {need_p:.2f} D {need_d:.2f}) | districts {n_d}")
print([(n, q, round(v[0] + v[1], 2)) for n, q, v in TOP])

# ---- slide on the outlook slide's frame (two chart boxes on the left, text on the right) ----
src = prs.slides[13]
sl = prs.slides.add_slide(src.slide_layout)
for ph in list(sl.shapes):
    ph._element.getparent().remove(ph._element)
boxes = []
for shp in src.shapes:
    if getattr(shp, "has_chart", False) and shp.has_chart:
        boxes.append((shp.left, shp.top, shp.width, shp.height))
        continue
    try:
        if shp.fill.type == 1 and str(shp.fill.fore_color.rgb) == "FFFF00":
            continue
    except (AttributeError, TypeError):
        pass
    el = copy.deepcopy(shp._element)
    for e in el.iter():
        for k, v in list(e.attrib.items()):
            if k in (RNS + "id", RNS + "embed", RNS + "link"):
                rel = src.part.rels[v]
                e.set(k, sl.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external else sl.part.relate_to(rel.target_part, rel.reltype))
    sl.shapes._spTree.append(el)
boxes.sort()
S = {s.name: s for s in sl.shapes}


def set_text(shape, text):
    ps = shape.text_frame._txBody.findall(qn("a:p"))
    rs = ps[0].findall(qn("a:r"))
    rs[0].find(qn("a:t")).text = text
    for r in rs[1:]:
        ps[0].remove(r)
    for p in ps[1:]:
        p.getparent().remove(p)


def chart(kind, box, cats, pet, die, vmax, step, cat_pt=9):
    cd = CategoryChartData(number_format="0.0")
    cd.categories = cats
    cd.add_series("Petrol", pet)
    cd.add_series("Diesel", die)
    gf = sl.shapes.add_chart(kind, box[0], box[1], box[2], box[3], cd)
    gf.name = "Chart"
    ch = gf.chart
    ch.has_title = False
    ch.font.size, ch.font.name = Pt(9), "Lato"
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.legend.font.size = Pt(8)
    for ax in (ch.category_axis, ch.value_axis):
        ax.has_major_gridlines = ax.has_minor_gridlines = False
        ax.major_tick_mark = ax.minor_tick_mark = XL_TICK_MARK.NONE
        ax.format.line.color.rgb = RGBColor.from_string("D4D4D4")
        ax.format.line.width = Emu(9525)
        ax.tick_labels.font.size = Pt(9)
        ax.tick_labels.font.color.rgb = RGBColor.from_string("666666")
    ch.category_axis.tick_labels.font.size = Pt(cat_pt)
    ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = "0", False
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 0, vmax, step
    pl = ch.plots[0]
    pl.gap_width, pl.overlap = 60, 100
    for ser, col in zip(pl.series, ("0A2373", "6782B4")):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
        ser.invert_if_negative = False
    for rc in ch._chartSpace.xpath("c:roundedCorners"):
        rc.set("val", "0")
    return ch


top_prov = max(sum(reg[c]) for c in codes)
chart(XL_CHART_TYPE.COLUMN_STACKED, boxes[0], [f"{n} {sum(reg[c]):.1f}" for c, n in MOI], [reg[c][0] for c, _ in MOI], [reg[c][1] for c, _ in MOI],
      2 * int(top_prov / 2 + 1), 2, cat_pt=8)
rev = TOP[::-1]
top_d = TOP[0][2][0] + TOP[0][2][1]
chart(XL_CHART_TYPE.BAR_STACKED, boxes[1], [f"{n} {v[0] + v[1]:.1f}" for n, _, v in rev], [v[0] for _, _, v in rev], [v[1] for _, _, v in rev],
      int(top_d) + 1, 1, cat_pt=8)
units = sorted([s for s in sl.shapes if s.name == "Evidence unit"], key=lambda s: (s.top, s.left))
assert len(units) == 3
set_text(units[0], "Billion litres, 2025 | 2022 district shares applied to 2025 national demand | totals beside each name")
set_text(units[1], "By province")
set_text(units[2], f"Largest magisterial districts (of {n_d})")
gk = sum(reg["GP"]) + sum(reg["KZN"])
set_text(S["Title 1"], f"The market of interest used about {mt:.1f}bn litres in 2025, of which about {need:.0f}bn litres had to be brought in")
set_text(S["Vopak exhibit left heading"], "Petrol and diesel demand in the market of interest, 2025")
set_text(S["Vopak exhibit right heading"], "From demand to import requirement")
set_text(S["Evidence insights"],
         f"Market of interest, 2025: {mt:.1f}bn L ({mt / nat:.0%} of national {nat:.1f}): petrol {mp:.1f}, diesel {md:.1f}. Gauteng and KwaZulu-Natal are {gk / mt:.0%}.\n\n"
         f"Made inside the area: Secunda and Natref, {inside:.1f}bn L (petrol {ins_p:.1f}, diesel {ins_d:.1f}; p10).\n\n"
         f"To be brought in: about {need:.1f}bn L, petrol {need_p:.1f} and diesel {need_d:.1f}, through Durban or other routes (p11).\n\n"
         f"Method: each district's 2022 share of national sales, by product, applied to 2025 national demand (p4). On 2022 sales alone the area was {mt22:.1f}bn L.")
set_text(S["Evidence takeaway"], f"Local refineries cover about {inside / mt:.0%} of the market of interest; the other {need / mt:.0%} must be imported or moved in.")
set_text(S["Unified source footer"],
         "Department fuel sales by magisterial district, four quarters of 2022 (latest full year), grown to 2025 at national growth by product (2025 demand = "
         "production + imports − exports, p4). Market of interest as in the AIA board report of November 2025: Gauteng, Limpopo, North West, Free State, "
         "Mpumalanga, KwaZulu-Natal. Production: p10. Assumes Secunda and Natref output stays in the area. Sales are recorded where fuel is sold.")
num = list(S["Slide Number Placeholder"]._element.iter(qn("a:t")))
assert len(num) == 1
num[0].text = "13"
prs.save(OUT)
import json                                                      # the funnel (p14) reads these, so the two slides agree
(Path(__file__).parent / "moi_2025.json").write_text(json.dumps({
    "national": nat25, "imports": {p: trade("import", p) for p in PD}, "moi": {"petrol": mp, "diesel": md},
    "inside": {"petrol": ins_p, "diesel": ins_d}, "need": {"petrol": need_p, "diesel": need_d},
    "kzn": {"petrol": reg["KZN"][0], "diesel": reg["KZN"][1]},
    "provinces": {c: {"name": n, "petrol": reg[c][0], "diesel": reg[c][1]} for c, n in MOI},
    "districts": [{"province": q, "district": n, "petrol": v[0], "diesel": v[1]} for n, q, v in md_list]}, indent=1), encoding="utf-8")
print("saved", OUT, len(prs.slides))
