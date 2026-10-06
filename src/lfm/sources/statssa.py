"""Statistics South Africa workbooks, read from disk.

Stats SA's website sits behind a bot check that refuses scripted downloads,
so these files cannot be fetched automatically. Once a year someone downloads
them in a browser into ``external/data/raw/statssa/`` and this module reads them:

    Gross domestic product (P0441), "GDP Time series" workbook
        statssa.gov.za > Publications > P0441 > latest quarter
        file name like "GDP P0441 - GDP Time series Q2 2026.xlsx"

The workbook's "Annual" sheet has one row per series. Column H03 is the
series code, H04 and H05 describe it, H15 gives the price basis, H17 the
unit, and the columns named Y1993, Y1994, ... hold the values.
"""
from __future__ import annotations

import re
from pathlib import Path

GDP_FILE_PATTERN = "GDP P0441*Time series*.xlsx"
ANNUAL_SHEET = "Annual"
GDP_CODE = "AR1000"            # GDP at market prices, constant prices, annual
VALUE_ADDED = "Value added at basic prices"

_YEAR_COLUMN = re.compile(r"^Y((?:19|20)\d{2})$")
_QUARTER = re.compile(r"Q([1-4])\s*(20\d{2})")


def latest_gdp_file(folder: Path) -> Path | None:
    """The GDP time-series workbook for the most recent quarter in ``folder``."""
    def key(path: Path) -> tuple[int, int]:
        match = _QUARTER.search(path.name)
        return (int(match.group(2)), int(match.group(1))) if match else (0, 0)

    files = [p for p in folder.glob(GDP_FILE_PATTERN) if not p.name.startswith("~$")]
    return max(files, key=key) if files else None


def read_annual_sheet(path: Path) -> list[list]:
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if ANNUAL_SHEET not in book.sheetnames:
            raise ValueError(f"{path.name}: no sheet named {ANNUAL_SHEET!r}")
        return [list(row) for row in book[ANNUAL_SHEET].iter_rows(values_only=True)]


def parse_constant_price_series(rows: list[list]) -> tuple[list[dict], list[str]]:
    """GDP and industry value added at constant prices, in rand.

    Returns ``(series, warnings)``. Each series is
    ``{"code", "name", "price_basis", "values": {year: rand}}``; ``name`` is
    ``"gdp"`` for total GDP and the industry description for value added.
    The workbook is in R million; values are converted to rand.
    """
    if not rows:
        return [], ["annual sheet is empty"]
    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    try:
        code_col, group_col, name_col = header.index("H03"), header.index("H04"), header.index("H05")
        basis_col, unit_col = header.index("H15"), header.index("H17")
    except ValueError:
        return [], ["annual sheet: expected column headings H03-H05, H15 and H17 not found"]
    years = {i: int(m.group(1)) for i, h in enumerate(header) if (m := _YEAR_COLUMN.match(h))}
    if not years:
        return [], ["annual sheet: no year columns found"]

    out: list[dict] = []
    warnings: list[str] = []
    for row in rows[1:]:
        if len(row) <= max(code_col, basis_col, unit_col):
            continue
        basis = str(row[basis_col] or "")
        if not basis.lower().startswith("constant"):
            continue
        code, group = str(row[code_col] or "").strip(), str(row[group_col] or "").strip()
        if code == GDP_CODE:
            name = "gdp"
        elif group == VALUE_ADDED:
            name = str(row[name_col] or "").strip()
        else:
            continue
        unit = str(row[unit_col] or "").strip().lower()
        if unit != "r million":
            warnings.append(f"{code}: unit is {row[unit_col]!r}, expected R million; skipped")
            continue
        values = {
            year: float(row[i]) * 1_000_000
            for i, year in years.items()
            if i < len(row) and isinstance(row[i], (int, float))
        }
        if values:
            out.append({"code": code, "name": name, "price_basis": basis, "values": values})
    if not any(s["name"] == "gdp" for s in out):
        warnings.append(f"GDP series {GDP_CODE} not found")
    return out, warnings


QUARTERLY_SHEET = "Quarterly"
QUARTERLY_GDP_CODE = "QRU1000"    # GDP at market prices, constant prices, not seasonally adjusted
_QUARTER_COLUMN = re.compile(r"^((?:19|20)\d{2})0([1-4])$")


def parse_quarterly_constant_price_series(rows: list[list]) -> tuple[list[dict], list[str]]:
    """Quarterly GDP and industry value added at constant prices, in rand.

    Takes the "Actual values" rows, not the seasonally adjusted and annualised
    ones, so the four quarters of a year add to the annual figure. Each series
    is ``{"code", "name", "price_basis", "values": {"2026-Q2": rand}}``.
    """
    if not rows:
        return [], ["quarterly sheet is empty"]
    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    try:
        code_col, group_col, name_col = header.index("H03"), header.index("H04"), header.index("H05")
        basis_col, kind_col, unit_col = header.index("H15"), header.index("H16"), header.index("H17")
    except ValueError:
        return [], ["quarterly sheet: expected column headings H03-H05 and H15-H17 not found"]
    quarters = {
        i: f"{m.group(1)}-Q{m.group(2)}"
        for i, h in enumerate(header) if (m := _QUARTER_COLUMN.match(h))
    }
    if not quarters:
        return [], ["quarterly sheet: no quarter columns found"]

    out: list[dict] = []
    warnings: list[str] = []
    for row in rows[1:]:
        if len(row) <= max(code_col, basis_col, kind_col, unit_col):
            continue
        basis = str(row[basis_col] or "")
        if not basis.lower().startswith("constant"):
            continue
        if str(row[kind_col] or "").strip().lower() != "actual values":
            continue
        code, group = str(row[code_col] or "").strip(), str(row[group_col] or "").strip()
        if code == QUARTERLY_GDP_CODE:
            name = "gdp"
        elif group == VALUE_ADDED:
            name = str(row[name_col] or "").strip()
        else:
            continue
        if str(row[unit_col] or "").strip().lower() != "r million":
            continue    # the "% of GDP" rows share the same headings
        values = {
            quarter: float(row[i]) * 1_000_000
            for i, quarter in quarters.items()
            if i < len(row) and isinstance(row[i], (int, float))
        }
        if values:
            out.append({"code": code, "name": name, "price_basis": basis, "values": values})
    if not any(s["name"] == "gdp" for s in out):
        warnings.append(f"quarterly GDP series {QUARTERLY_GDP_CODE} not found")
    return out, warnings


# --------------------------------------------------------------------------- #
# Monthly releases in Stats SA's time-series layout
#
#     statssa.gov.za > Time series data > Excel
#         "P2041 Mining Production and sales(202607).zip"
#         "P3041.2 Manufacturing_ Production and sales(202607).zip"
#         "P7162 Land transport survey(202607).zip"
#         "P0141 - CPI(COICOP) from Jan 2008 (202608).zip"
#
# Each zip holds one or more workbooks with a single sheet: descriptor columns
# headed H01, H02, ... and one column per month headed MO<mm><yyyy>. H03 is the
# series code. The zips are kept as downloaded and read in place.

_MONTH_COLUMN = re.compile(r"^MO(0[1-9]|1[0-2])((?:19|20)\d{2})$")
_RELEASE_STAMP = re.compile(r"\((20\d{4})\)")

# release -> (file pattern, workbook inside the zip, {series code: (series name, unit)})
MONTHLY_RELEASES: dict[str, tuple[str, str, dict[str, tuple[str, str]]]] = {
    "mining": ("P2041 Mining Production and sales(*).zip", "from 2003", {
        "FMP20000": ("mining_volume_total", "index, 2019=100"),
        "FMP20001": ("mining_volume_excluding_gold", "index, 2019=100"),
        "FMP21000": ("mining_volume_coal", "index, 2019=100"),
    }),
    "manufacturing": ("P3041.2 Manufacturing*Production and sales(*).zip", "from 1998", {
        "MPI30000": ("manufacturing_volume_total", "index, 2019=100"),
    }),
    "land_transport": ("P7162 Land transport survey(*).zip", "Land transport", {
        "payl_totl": ("freight_payload_total", "thousand tonnes"),
        "roadpayl": ("freight_payload_road", "thousand tonnes"),
        "railpayl": ("freight_payload_rail", "thousand tonnes"),
        "nops_totl": ("passenger_journeys_total", "thousand journeys"),
        "roadnops": ("passenger_journeys_road", "thousand journeys"),
        "railnops": ("passenger_journeys_rail", "thousand journeys"),
    }),
    "consumer_prices": ("P0141 - CPI(COICOP) from Jan 2008 (*).zip", "CPI", {
        "CPI60001": ("consumer_price_index_headline", "index, December 2024=100"),
    }),
}


def latest_release_file(folder: Path, pattern: str) -> Path | None:
    """The newest zip matching ``pattern``, by the (yyyymm) stamp in its name."""
    def key(path: Path) -> str:
        match = _RELEASE_STAMP.search(path.name)
        return match.group(1) if match else ""

    files = [p for p in folder.glob(pattern) if key(p)]
    return max(files, key=key) if files else None


def read_release_zip(path: Path, member_contains: str) -> list[list]:
    """Rows of the workbook inside a release zip whose name contains ``member_contains``."""
    import io
    import warnings
    import zipfile

    import openpyxl

    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist()
                 if n.lower().endswith(".xlsx") and member_contains.lower() in n.lower()]
        if not names:
            raise ValueError(f"{path.name}: no workbook with {member_contains!r} in its name")
        data = archive.read(names[0])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sheet = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if row[0] is None:      # the sheets declare a million rows; data stops at the first blank
                break
            rows.append(list(row))
        return rows


def parse_monthly_series(rows: list[list], wanted: dict[str, tuple[str, str]]) -> tuple[list[dict], list[str]]:
    """Monthly values for the series codes in ``wanted``.

    Returns ``(series, warnings)``; each series is
    ``{"code", "name", "unit", "values": {"2026-07": float}}``. A wanted code
    that is not in the sheet is reported, never silently skipped.
    """
    if not rows:
        return [], ["sheet is empty"]
    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if "H03" not in header:
        return [], ["expected column heading H03 (series code) not found"]
    code_col = header.index("H03")
    months = {
        i: f"{m.group(2)}-{m.group(1)}" for i, h in enumerate(header) if (m := _MONTH_COLUMN.match(h))
    }
    if not months:
        return [], ["no month columns found"]
    out: list[dict] = []
    found: set[str] = set()
    for row in rows[1:]:
        code = str(row[code_col] or "").strip()
        if code not in wanted or code in found:
            continue
        found.add(code)
        values = {}
        for i, month in months.items():
            cell = row[i] if i < len(row) else None
            if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                values[month] = float(cell)
            elif isinstance(cell, str) and cell.strip():
                try:
                    values[month] = float(cell.replace(",", ".").replace(" ", ""))
                except ValueError:
                    pass                      # ".." and similar mark a month not available
        name, unit = wanted[code]
        out.append({"code": code, "name": name, "unit": unit, "values": values})
    missing = sorted(set(wanted) - found)
    warnings = [f"series {code} not found" for code in missing]
    return out, warnings


# --------------------------------------------------------------------------- #
# Provincial GDP (P0441.2), annual
#
#     statssa.gov.za > Time series data > Excel
#         "P0441.2  Provincial Gross Domestic Product(2024).zip"
#
# One workbook. Tables 2 to 10 are one province each, titled
# "<Province> – GDPR by activity", and hold four blocks down the sheet:
# a. current prices, b. percentage contributions, c. constant 2015 prices,
# d. percentage changes. Block c is read: a row of years, then one row per
# industry down to "GDPR at market prices". Values are in R million.

PROVINCIAL_GDP_PATTERN = "P0441.2*Provincial Gross Domestic Product(*).zip"
PROVINCE_CODES = {
    "western cape": "WC", "eastern cape": "EC", "northern cape": "NC", "free state": "FS",
    "kwazulu-natal": "KZN", "north west": "NW", "gauteng": "GP", "mpumalanga": "MP", "limpopo": "LP",
}
_PROVINCE_TITLE = re.compile(r"^\s*(.+?)\s*[–-]\s*GDPR by activity", re.IGNORECASE)
_STAMP_YEAR = re.compile(r"\(((?:19|20)\d{2})\)")


def latest_provincial_gdp_file(folder: Path) -> Path | None:
    def key(path: Path) -> str:
        match = _STAMP_YEAR.search(path.name)
        return match.group(1) if match else ""

    files = [p for p in folder.glob(PROVINCIAL_GDP_PATTERN) if key(p)]
    return max(files, key=key) if files else None


def read_zip_sheets(path: Path) -> dict[str, list[list]]:
    """Every sheet of the first workbook inside a release zip, as rows."""
    import io
    import warnings
    import zipfile

    import openpyxl

    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith(".xlsx")]
        if not names:
            raise ValueError(f"{path.name}: no workbook inside")
        data = archive.read(names[0])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in book}


def parse_provincial_gdp(sheets: dict[str, list[list]]) -> tuple[list[dict], list[str]]:
    """Constant-price value added and GDP by province and industry, in rand.

    Returns ``(rows, warnings)``; each row is
    ``{"province", "industry", "year", "value"}``. A province sheet whose
    constant-price block cannot be found is reported and left out, and so is
    any of the nine provinces with no sheet.
    """
    out: list[dict] = []
    warnings: list[str] = []
    seen: set[str] = set()
    for name, rows in sheets.items():
        title = next((str(c) for c in (rows[0] if rows else []) if c is not None), "")
        match = _PROVINCE_TITLE.match(title)
        if not match:
            continue
        province = PROVINCE_CODES.get(match.group(1).strip().lower())
        if province is None:
            warnings.append(f"{name}: province {match.group(1)!r} not recognised")
            continue
        start = next((i for i, row in enumerate(rows)
                      if isinstance(row[0], str) and row[0].strip().lower().startswith("c. constant")), None)
        if start is None or start + 1 >= len(rows):
            warnings.append(f"{name}: constant-price block not found")
            continue
        header = rows[start + 1]
        years = {}
        for i, cell in enumerate(header):
            text = str(cell).strip() if cell is not None else ""
            if re.fullmatch(r"(?:19|20)\d{2}", text):
                years[i] = int(text)
        if not years:
            warnings.append(f"{name}: no year columns in the constant-price block")
            continue
        seen.add(province)
        for row in rows[start + 2:]:
            label = row[0]
            if not isinstance(label, str) or not label.strip():
                break
            if label.strip().lower().startswith("d."):
                break
            for i, year in years.items():
                cell = row[i] if i < len(row) else None
                if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                    out.append({"province": province, "industry": label.strip(), "year": year,
                                "value": float(cell) * 1_000_000})
    for code in PROVINCE_CODES.values():
        if code not in seen:
            warnings.append(f"no sheet read for province {code}")
    return out, warnings


# --------------------------------------------------------------------------- #
# Mid-year population estimates (P0302)
#
#     statssa.gov.za > Publications > P0302 > latest release
#     "Country projection by population group, sex and age (2002-2026).xlsx"
#         one row per population group x sex x age band, one column per year
#     "MYPE report table website_ 2026.xlsx"
#         the release's summary tables; its national total is used as a check

POPULATION_FILE_PATTERN = "Country projection by population group*.xlsx"
SUMMARY_FILE_PATTERN = "MYPE report table*.xlsx"
SUMMARY_SHEET = "MYPE by pop grp and sex"

_END_YEAR = re.compile(r"\((?:19|20)\d{2}\s*-\s*((?:19|20)\d{2})\)")
_RELEASE_YEAR = re.compile(r"((?:19|20)\d{2})")


def latest_population_file(folder: Path) -> Path | None:
    """The country projection workbook whose year range ends latest."""
    def key(path: Path) -> int:
        match = _END_YEAR.search(path.name)
        return int(match.group(1)) if match else 0

    files = [p for p in folder.glob(POPULATION_FILE_PATTERN) if not p.name.startswith("~$")]
    return max(files, key=key) if files else None


def latest_summary_file(folder: Path) -> Path | None:
    def key(path: Path) -> int:
        years = _RELEASE_YEAR.findall(path.stem)
        return int(years[-1]) if years else 0

    files = [p for p in folder.glob(SUMMARY_FILE_PATTERN) if not p.name.startswith("~$")]
    return max(files, key=key) if files else None


def read_sheet(path: Path, sheet: str | None = None) -> list[list]:
    """Rows of one sheet (the first if ``sheet`` is not given)."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet is not None and sheet not in book.sheetnames:
            raise ValueError(f"{path.name}: no sheet named {sheet!r}")
        ws = book[sheet] if sheet is not None else book[book.sheetnames[0]]
        return [list(row) for row in ws.iter_rows(values_only=True)]


def parse_population(rows: list[list]) -> tuple[dict[int, float], list[str]]:
    """Total population by year: ``({year: persons}, warnings)``.

    Adds up every population group, sex and age band. Rows that are already
    totals are left out so nothing is counted twice.
    """
    header_index = next(
        (i for i, row in enumerate(rows)
         if sum(1 for c in row if _is_year(c)) >= 5 and any(_text(c) == "sex" for c in row)),
        None,
    )
    if header_index is None:
        return {}, ["population sheet: heading row with Sex and year columns not found"]
    header = rows[header_index]
    # The sheet carries further tables to the right (by sex and age, then by
    # age only) that repeat the same years. Only the first, most detailed
    # block is read: the label columns and the unbroken run of years after them.
    first_year = next(i for i, c in enumerate(header) if _is_year(c))
    years: dict[int, int] = {}
    for i in range(first_year, len(header)):
        if not _is_year(header[i]):
            break
        years[i] = int(float(header[i]))
    labels = [i for i in range(first_year) if isinstance(header[i], str)]

    totals = dict.fromkeys(years.values(), 0.0)
    counted = 0
    for row in rows[header_index + 1:]:
        words = [_text(row[i]) for i in labels if i < len(row)]
        if not any(words) or any("total" in w or w == "all" for w in words):
            continue
        values = {year: row[i] for i, year in years.items() if i < len(row)}
        if not all(isinstance(v, (int, float)) for v in values.values()):
            continue
        for year, value in values.items():
            totals[year] += float(value)
        counted += 1
    if counted == 0:
        return {}, ["population sheet: no data rows found"]
    return totals, []


def parse_summary_total(rows: list[list]) -> float | None:
    """The national total printed in the release's summary table, if found."""
    for row in rows:
        if row and _text(row[0]) == "total":
            numbers = [c for c in row if isinstance(c, (int, float)) and c > 1_000_000]
            if numbers:
                return float(numbers[-1])
    return None


def average_growth(population: dict[int, float], years: int) -> tuple[float, int, int]:
    """Average annual growth over the last ``years`` years: ``(rate, from, to)``."""
    last = max(population)
    first = last - years
    if first not in population:
        raise ValueError(f"population series does not reach back to {first}")
    return (population[last] / population[first]) ** (1 / years) - 1, first, last


def _is_year(cell) -> bool:
    try:
        return 1900 <= int(float(str(cell).strip())) <= 2100 and len(str(cell).strip()) in (4, 6)
    except (TypeError, ValueError):
        return False


def _text(cell) -> str:
    return str(cell).strip().lower() if isinstance(cell, str) else ""
