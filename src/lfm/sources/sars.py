"""Customs trade in refined fuels from the SARS trade statistics portal.

The portal's "Trade Data Download" page is a web form, not a file listing. It
has to be walked in the order a person would: trade type, focus area, chapter,
year, then Download. The page then answers with a small script that opens
``Download.aspx`` in the same session, and that address returns the workbook.

Things the form insists on, found by trial:
    * every country must be ticked, even when the focus is tariff lines;
    * at most two years per download (one is used here);
    * the tariff list only appears after the chapter has been posted back.

If the download stops working, check those three first, then the control names
below against the page source.

Quantities are in the unit SARS records for the line: litres from 2014 for the
fuels here, kilograms before that (2013 is a mix). They are never converted.
"""
from __future__ import annotations

import http.cookiejar
import re
import urllib.parse
import urllib.request
from pathlib import Path

URL = "https://tools.sars.gov.za/tradestatsportal/data_download.aspx"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"
CHAPTER = "27"            # crude, coal, petroleum and electricity
FIRST_YEAR = 2010         # earliest year the portal offers
TRADE_TYPES = ("Imports", "Exports")
_PREFIX = "ctl00$ContentPlaceHolder1$"

# Tariff lines by product. The same product sits under 2710.11 (to 2011),
# 2710.12 (from 2012) and, for a few entries, 2710.19.
PRODUCT_TARIFFS: dict[str, tuple[str, ...]] = {
    "petrol": ("27101102", "27101202"),
    "diesel": ("27101130", "27101230", "27101930"),
    "diesel_biodiesel_blend": ("27102000",),
    "jet": ("27101107", "27101207", "27101907"),
    "paraffin": ("27101115", "27101126", "27101215", "27101226", "27101915", "27101926"),
    "fuel_oil": ("27101135", "27101235", "27101935"),
}
TARIFF_PRODUCT = {code: product for product, codes in PRODUCT_TARIFFS.items() for code in codes}
UNITS = {"LI": "litres", "KG": "kilograms"}


# --------------------------------------------------------------------------- #
# Reading the form

def hidden_fields(html: str) -> dict[str, str]:
    """ASP.NET state fields (``__VIEWSTATE`` and friends) to send back."""
    return {
        m.group(1): m.group(2)
        for m in re.finditer(
            r'<input type="hidden" name="(__[A-Z]+)" id="[^"]*" value="([^"]*)"', html)
    }


def checkbox_labels(html: str, control: str) -> list[tuple[int, str]]:
    """``(index, label)`` for each box of a tick-list such as ``ddlTariffs``."""
    pattern = (r'<input id="ctl00_ContentPlaceHolder1_%s_(\d+)"[^>]*/>\s*<label[^>]*>([^<]*)</label>'
               % re.escape(control))
    return [(int(m.group(1)), m.group(2).strip()) for m in re.finditer(pattern, html)]


def download_link(html: str) -> str | None:
    """The address the page opens once it has accepted the selection."""
    match = re.search(r"window\.open\('(Download\.aspx[^']*)'", html)
    return urllib.parse.urljoin(URL, match.group(1)) if match else None


def refusal(html: str) -> str | None:
    """The message the page shows when it rejects a selection, if any."""
    match = re.search(r"alert\('([^']*)'\)", html)
    return match.group(1) if match else None


def download_year(trade_type: str, year: int, timeout: int = 600) -> bytes:
    """One year of imports or exports for the fuel tariff lines, as a workbook."""
    if trade_type not in TRADE_TYPES:
        raise ValueError(f"trade type must be one of {TRADE_TYPES}")
    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    opener.addheaders = [("User-Agent", USER_AGENT)]

    def post(html: str, target: str, state: dict[str, str], extra: dict | None = None) -> str:
        data = hidden_fields(html) | state
        data |= {"__EVENTTARGET": target, "__EVENTARGUMENT": "", "__LASTFOCUS": ""}
        data |= extra or {}
        body = urllib.parse.urlencode(data).encode()
        with opener.open(URL, body, timeout=timeout) as response:
            return response.read().decode("utf-8", "ignore")

    html = opener.open(URL, timeout=timeout).read().decode("utf-8", "ignore")
    years = {label: index for index, label in checkbox_labels(html, "ddlYears")}
    if str(year) not in years:
        raise ValueError(f"the portal does not offer {year}")

    state = {_PREFIX + "tradeType": "rdb" + trade_type}
    html = post(html, _PREFIX + "rdb" + trade_type, state)
    state[_PREFIX + "focusArea"] = "rdbTariffs"
    html = post(html, _PREFIX + "rdbTariffs", state)

    chapters = [i for i, label in checkbox_labels(html, "ddlChapters") if label.startswith(CHAPTER)]
    if not chapters:
        raise ValueError(f"chapter {CHAPTER} not found on the form")
    state[f"{_PREFIX}ddlChapters${chapters[0]}"] = "on"
    html = post(html, _PREFIX + "ddlChapters", state)

    wanted = [i for i, label in checkbox_labels(html, "ddlTariffs") if label[:8] in TARIFF_PRODUCT]
    if not wanted:
        raise ValueError("no fuel tariff lines found on the form")
    for index in wanted:
        state[f"{_PREFIX}ddlTariffs${index}"] = "on"
    state[f"{_PREFIX}ddlYears${years[str(year)]}"] = "on"
    html = post(html, _PREFIX + "ddlYears", state)

    for index, _ in checkbox_labels(html, "ddlMonths"):
        state[f"{_PREFIX}ddlMonths${index}"] = "on"
    for index, _ in checkbox_labels(html, "chkColumns"):
        state[f"{_PREFIX}chkColumns${index}"] = "on"
    for index, _ in checkbox_labels(html, "ddlCountries"):
        state[f"{_PREFIX}ddlCountries${index}"] = "on"
    state |= {_PREFIX + "chkAll": "on", _PREFIX + "hdnTariff": "visble",
              _PREFIX + "hdnMonth": "visble"}
    html = post(html, "", state, {_PREFIX + "btnDownload": "Download"})

    link = download_link(html)
    if link is None:
        raise ValueError(f"the portal did not offer a file: {refusal(html) or 'no reason given'}")
    with opener.open(link, timeout=timeout) as response:
        data = response.read()
    if data[:2] != b"PK":
        raise ValueError("the portal's download was not a workbook")
    return data


# --------------------------------------------------------------------------- #
# Reading a downloaded workbook

COLUMNS = ("TradeType", "DistrictOfficeName", "CountryOfOriginName", "CountryOfDestinationName",
           "Tariff", "StatisticalUnit", "TransportCodeDescription", "YearMonth",
           "StatisticalQuantity", "CustomsValue")


def read_report(path: Path) -> list[dict]:
    """Rows of a downloaded report, one per tariff line, office, partner, mode and month."""
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(rows)]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise ValueError(f"{path.name}: columns missing: {missing}")
    at = {name: header.index(name) for name in COLUMNS}
    out = []
    for row in rows:
        if row[at["Tariff"]] is None:
            continue
        tariff = str(int(row[at["Tariff"]]))
        flow = "import" if str(row[at["TradeType"]]).lower().startswith("import") else "export"
        year_month = str(int(row[at["YearMonth"]]))
        out.append({
            "flow": flow,
            "product": TARIFF_PRODUCT.get(tariff, "other"),
            "tariff": tariff,
            "period": f"{year_month[:4]}-{year_month[4:]}",
            "district_office": str(row[at["DistrictOfficeName"]] or "").strip(),
            "transport_mode": str(row[at["TransportCodeDescription"]] or "").strip(),
            "partner": str(row[at["CountryOfOriginName" if flow == "import"
                                   else "CountryOfDestinationName"]] or "").strip(),
            "unit": _unit(row[at["StatisticalUnit"]]),
            "quantity": float(row[at["StatisticalQuantity"]] or 0.0),
            "customs_value_rand": float(row[at["CustomsValue"]] or 0.0),
        })
    return out


def _unit(cell) -> str:
    code = str(cell).strip() if cell is not None else ""
    return UNITS.get(code, code or "not stated")


def annual(rows: list[dict], by: tuple[str, ...] = ()) -> list[dict]:
    """Yearly totals by flow, product and unit, and by any extra keys in ``by``.

    ``months_reported`` is the number of distinct months behind each total, so
    a part year (the current one) is visible as such.
    """
    keys = ("flow", "product", "unit") + by
    totals: dict[tuple, list] = {}
    for row in rows:
        key = (int(row["period"][:4]),) + tuple(row[k] for k in keys)
        entry = totals.setdefault(key, [0.0, set()])
        entry[0] += row["quantity"]
        entry[1].add(row["period"])
    return [
        {"period": key[0], **dict(zip(keys, key[1:])), "value": round(total, 3),
         "months_reported": len(months)}
        for key, (total, months) in sorted(totals.items())
    ]
