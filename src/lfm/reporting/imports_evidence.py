"""Workbook import evidence and the existing Python domestic-gap calculation.

Workbook cached forecasts are comparison evidence, not available-import inputs.
No workbook file enters the calculation engine.
"""
import openpyxl

from lfm.assumptions import YamlDirectoryProvider
from lfm.run import Run
from lfm.scripts.compare_history import model_historical
from lfm.model.supply.flows import compute_balance, compute_supply


def imports_evidence(workbook):
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    sheet = wb['fImports']
    records = []
    for label, product, start in [('Petrol', 'gasoline', 82), ('Diesel', 'diesel', 46), ('Jet', 'jet_a1', 10)]:
        history_row, forecast_row = start + 12, start + 20
        assert sheet.cell(history_row, 1).value == 2022
        assert sheet.cell(forecast_row, 1).value == 2030
        records.append(dict(fuel=label, product=product,
                            historical=sheet.cell(history_row, 4).value,
                            excel_h=sheet.cell(forecast_row, 13).value,
                            excel_l=sheet.cell(forecast_row, 22).value,
                            cells=f'D{history_row}; M{forecast_row}; V{forecast_row}'))
    wb.close()
    provider = YamlDirectoryProvider()
    for scenario in ('high_demand', 'low_demand'):
        run = Run(vintage='2026', scenario=scenario)
        wide = model_historical(run, provider)
        long = wide.melt(id_vars=['country', 'period'], var_name='product', value_name='volume').dropna(subset=['volume'])
        balance = compute_balance(long, compute_supply(provider, run))
        balance = balance.query("country == 'ZAF' and period == 2030").set_index('refinery_product')
        for record in records:
            record[scenario] = float(balance.loc[record['product'], 'deficit_litres'])
    return records
